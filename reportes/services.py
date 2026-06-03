"""
Servicio de reportes y exportación (Hito 8).

Genera reportes aptos para inspección del Ministerio de Trabajo y auditoría
BASC, además de la exportación individual de datos personales (Habeas Data,
Ley 1581). Cada reporte se entrega en Excel (openpyxl) o PDF (reportlab).

ARQUITECTURA (tres capas separadas, a propósito):
  1. RECOLECCIÓN  -> ``recolectar_*`` devuelven una estructura ``Reporte`` con
     tablas/filas. Es solo lectura sobre datos existentes (no toca marcaciones
     ni audit_log: no rompe inmutabilidad ni triggers).
  2. RENDER       -> ``render_excel`` / ``render_pdf`` convierten un ``Reporte``
     en bytes. No saben nada del dominio.
  3. ORQUESTACIÓN -> ``exportar`` renderiza Y registra SIEMPRE la exportación en
     audit_log (accion=EXPORTACION_DATOS). Ningún camino (API o admin) puede
     materializar un archivo sin dejar rastro: una exportación es un acceso
     masivo a datos personales y debe trazarse.

CUMPLIMIENTO — DATOS SENSIBLES DE SALUD (Ley 1581):
  Los soportes de novedades (p. ej. incapacidades médicas) NUNCA se vuelcan en
  un reporte. Solo se REFERENCIAN sus metadatos (nombre, tipo, tamaño, fecha).
  El archivo médico jamás se lee ni se incrusta.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal
from io import BytesIO
from typing import Any, Iterable, Optional
from xml.sax.saxutils import escape

from django.db.models import Q
from django.http import HttpResponse
from django.utils import timezone

from auditoria.models import AuditLog
from auditoria.services import AuditService

XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
PDF_CONTENT_TYPE = "application/pdf"

# Aviso fijo que acompaña cualquier referencia a un soporte sensible.
NOTA_SOPORTE_SENSIBLE = "archivo NO incluido (dato sensible, Ley 1581)"


# ==========================================================================
# Estructura genérica de un reporte
# ==========================================================================
@dataclass
class Tabla:
    nombre: str
    columnas: list[str]
    filas: list[list[Any]] = field(default_factory=list)


@dataclass
class Reporte:
    titulo: str
    subtitulo: str
    meta: dict[str, Any]
    tablas: list[Tabla] = field(default_factory=list)

    def iter_celdas(self) -> Iterable[str]:
        """Itera el texto de TODAS las celdas (útil para verificaciones)."""
        for tabla in self.tablas:
            for fila in tabla.filas:
                for celda in fila:
                    yield "" if celda is None else str(celda)

    def total_filas(self) -> int:
        return sum(len(t.filas) for t in self.tablas)


@dataclass
class ExportResult:
    contenido: bytes
    filename: str
    content_type: str


# ==========================================================================
# Utilidades de formato
# ==========================================================================
def _hora_local(dt) -> str:
    if not dt:
        return ""
    return timezone.localtime(dt).strftime("%Y-%m-%d %H:%M")


def _si_no(valor) -> str:
    return "Sí" if valor else "No"


def _num(valor) -> float:
    """Normaliza Decimals/None a float para escritura numérica en Excel."""
    if valor is None:
        return 0.0
    if isinstance(valor, Decimal):
        return float(valor)
    return float(valor)


def _hora_oficial(marcacion):
    """La hora que cuenta para asistencia: la del QR firmado, si no la del servidor."""
    return marcacion.timestamp_qr or marcacion.timestamp_servidor


def _subtitulo(meta: dict) -> str:
    partes = []
    if meta.get("Periodo"):
        partes.append(f"Periodo: {meta['Periodo']}")
    if meta.get("Sede"):
        partes.append(f"Sede: {meta['Sede']}")
    return " · ".join(partes)


# ==========================================================================
# RENDER — Excel
# ==========================================================================
def _sheet_title(nombre: str, usados: set) -> str:
    invalidos = set('[]:*?/\\')
    base = "".join("_" if c in invalidos else c for c in nombre)[:31] or "Hoja"
    titulo = base
    i = 1
    while titulo in usados:
        i += 1
        sufijo = f"_{i}"
        titulo = base[: 31 - len(sufijo)] + sufijo
    usados.add(titulo)
    return titulo


def render_excel(reporte: Reporte) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    portada = wb.active
    portada.title = "Portada"
    usados = {"Portada"}

    portada["A1"] = reporte.titulo
    portada["A1"].font = Font(bold=True, size=14)
    portada["A2"] = reporte.subtitulo
    fila = 4
    for clave, valor in reporte.meta.items():
        portada.cell(row=fila, column=1, value=str(clave)).font = Font(bold=True)
        portada.cell(row=fila, column=2, value=str(valor))
        fila += 1
    portada.column_dimensions["A"].width = 26
    portada.column_dimensions["B"].width = 50

    encabezado_fill = PatternFill("solid", fgColor="DDDDDD")
    for tabla in reporte.tablas:
        ws = wb.create_sheet(_sheet_title(tabla.nombre, usados))
        ws.cell(row=1, column=1, value=tabla.nombre).font = Font(bold=True, size=12)
        for ci, col in enumerate(tabla.columnas, start=1):
            celda = ws.cell(row=2, column=ci, value=str(col))
            celda.font = Font(bold=True)
            celda.fill = encabezado_fill
        for ri, fila_datos in enumerate(tabla.filas, start=3):
            for ci, valor in enumerate(fila_datos, start=1):
                if isinstance(valor, (int, float)):
                    ws.cell(row=ri, column=ci, value=valor)
                else:
                    ws.cell(row=ri, column=ci, value="" if valor is None else str(valor))
        # Ancho aproximado por columna.
        for ci, col in enumerate(tabla.columnas, start=1):
            ancho = len(str(col))
            for fila_datos in tabla.filas:
                if ci - 1 < len(fila_datos):
                    ancho = max(ancho, len(str(fila_datos[ci - 1])))
            ws.column_dimensions[get_column_letter(ci)].width = min(max(ancho + 2, 10), 60)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ==========================================================================
# RENDER — PDF
# ==========================================================================
def render_pdf(reporte: Reporte) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=1.2 * cm,
        rightMargin=1.2 * cm,
        topMargin=1.2 * cm,
        bottomMargin=1.2 * cm,
        title=reporte.titulo,
    )
    estilos = getSampleStyleSheet()
    estilo_celda = ParagraphStyle(
        "celda", parent=estilos["Normal"], fontSize=7, leading=8
    )
    estilo_encabezado = ParagraphStyle(
        "encabezado",
        parent=estilo_celda,
        fontName="Helvetica-Bold",
        textColor=colors.white,
    )

    historia = [Paragraph(escape(reporte.titulo), estilos["Title"])]
    if reporte.subtitulo:
        historia.append(Paragraph(escape(reporte.subtitulo), estilos["Heading3"]))
    meta_linea = " · ".join(
        f"<b>{escape(str(k))}:</b> {escape(str(v))}" for k, v in reporte.meta.items()
    )
    if meta_linea:
        historia.append(Paragraph(meta_linea, estilo_celda))
    historia.append(Spacer(1, 0.4 * cm))

    for tabla in reporte.tablas:
        historia.append(Paragraph(escape(tabla.nombre), estilos["Heading2"]))
        if not tabla.filas:
            historia.append(Paragraph("Sin registros en el periodo.", estilo_celda))
            historia.append(Spacer(1, 0.3 * cm))
            continue
        ncols = len(tabla.columnas)
        ancho_col = [doc.width / ncols] * ncols  # fuerza que la tabla quepa
        datos = [
            [Paragraph(escape(str(c)), estilo_encabezado) for c in tabla.columnas]
        ]
        for fila_datos in tabla.filas:
            datos.append(
                [
                    Paragraph(escape("" if v is None else str(v)), estilo_celda)
                    for v in fila_datos
                ]
            )
        flowable = Table(datos, colWidths=ancho_col, repeatRows=1)
        flowable.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0EA5E9")),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor("#F1F5F9")],
                    ),
                ]
            )
        )
        historia.append(flowable)
        historia.append(Spacer(1, 0.4 * cm))

    doc.build(historia)
    return buffer.getvalue()


# ==========================================================================
# RECOLECCIÓN — reportes agregados
# ==========================================================================
def recolectar_asistencia(marcaciones_qs, *, meta: dict) -> Reporte:
    """Marcaciones del periodo con su HORA OFICIAL, marcando offline / fuera de sede."""
    qs = marcaciones_qs.select_related("empleado", "sede").order_by(
        "empleado__apellidos", "empleado__nombres", "minuto_marcacion"
    )
    columnas = [
        "Documento",
        "Empleado",
        "Sede",
        "Tipo",
        "Hora oficial",
        "Origen",
        "Fuera de sede",
        "Distancia (m)",
        "Corrección",
    ]
    filas = [
        [
            m.empleado.documento_identidad,
            m.empleado.nombre_completo,
            m.sede.nombre,
            m.get_tipo_display(),
            _hora_local(_hora_oficial(m)),
            "Offline" if m.es_offline else "Online",
            _si_no(m.fuera_de_sede),
            _num(m.distancia_metros),
            _si_no(m.corrige_a_id is not None),
        ]
        for m in qs
    ]
    return Reporte(
        titulo="Reporte de asistencia",
        subtitulo=_subtitulo(meta),
        meta=meta,
        tablas=[Tabla("Asistencia", columnas, filas)],
    )


def recolectar_excepciones(marcaciones_qs, *, meta: dict) -> Reporte:
    """Marcaciones fuera de geocerca o marcadas para revisión offline (control BASC)."""
    qs = (
        marcaciones_qs.filter(Q(fuera_de_sede=True) | Q(revisar_offline=True))
        .select_related("empleado", "sede")
        .order_by("empleado__apellidos", "minuto_marcacion")
    )
    columnas = [
        "Documento",
        "Empleado",
        "Sede",
        "Tipo",
        "Hora oficial",
        "Fuera de sede",
        "Distancia (m)",
        "Offline",
        "Revisar offline",
    ]
    filas = [
        [
            m.empleado.documento_identidad,
            m.empleado.nombre_completo,
            m.sede.nombre,
            m.get_tipo_display(),
            _hora_local(_hora_oficial(m)),
            _si_no(m.fuera_de_sede),
            _num(m.distancia_metros),
            _si_no(m.es_offline),
            _si_no(m.revisar_offline),
        ]
        for m in qs
    ]
    return Reporte(
        titulo="Reporte de excepciones (control BASC)",
        subtitulo=_subtitulo(meta),
        meta=meta,
        tablas=[Tabla("Excepciones", columnas, filas)],
    )


def recolectar_horas(resumenes_qs, *, meta: dict) -> Reporte:
    """Horas clasificadas (ordinarias, extra diurnas/nocturnas, recargo, dom/festivo)."""
    qs = resumenes_qs.select_related("empleado", "empleado__sede").order_by(
        "empleado__apellidos", "fecha"
    )
    columnas = [
        "Documento",
        "Empleado",
        "Sede",
        "Fecha",
        "H. Ordinarias",
        "Extra Diurnas",
        "Extra Nocturnas",
        "Recargo Nocturno",
        "Dominical/Festivo",
        "Tardanza (min)",
        "Salida temprana (min)",
    ]
    filas = []
    tot = [0.0, 0.0, 0.0, 0.0, 0.0]
    for r in qs:
        valores = [
            _num(r.horas_ordinarias),
            _num(r.horas_extra_diurnas),
            _num(r.horas_extra_nocturnas),
            _num(r.horas_recargo_nocturno),
            _num(r.horas_dominical_festivo),
        ]
        for i, v in enumerate(valores):
            tot[i] += v
        filas.append(
            [
                r.empleado.documento_identidad,
                r.empleado.nombre_completo,
                r.empleado.sede.nombre,
                r.fecha.isoformat(),
                *valores,
                r.tardanza_minutos,
                r.salida_temprana_minutos,
            ]
        )
    if filas:
        filas.append(
            ["", "", "", "TOTALES", *[round(x, 2) for x in tot], "", ""]
        )
    return Reporte(
        titulo="Reporte de horas (clasificación legal)",
        subtitulo=_subtitulo(meta),
        meta=meta,
        tablas=[Tabla("Horas", columnas, filas)],
    )


def recolectar_novedades(novedades_qs, *, meta: dict) -> Reporte:
    """Novedades por tipo/estado. NO vuelca el contenido de soportes médicos."""
    from novedades.models import Novedad

    qs = (
        novedades_qs.select_related("empleado", "empleado__sede")
        .prefetch_related("soportes")
        .order_by("tipo", "estado", "fecha_inicio")
    )

    # Tabla 1: resumen por tipo × estado.
    conteo: dict[tuple, int] = {}
    detalle = []
    for n in qs:
        clave = (n.get_tipo_display(), n.get_estado_display())
        conteo[clave] = conteo.get(clave, 0) + 1
        soportes = list(n.soportes.all())
        detalle.append(
            [
                n.empleado.documento_identidad,
                n.empleado.nombre_completo,
                n.empleado.sede.nombre,
                n.get_tipo_display(),
                n.get_estado_display(),
                n.fecha_inicio.isoformat(),
                n.fecha_fin.isoformat() if n.fecha_fin else "",
                _si_no(bool(soportes)),
                len(soportes),
            ]
        )

    resumen_filas = [
        [tipo, estado, cant] for (tipo, estado), cant in sorted(conteo.items())
    ]
    tabla_resumen = Tabla(
        "Resumen por tipo y estado", ["Tipo", "Estado", "Cantidad"], resumen_filas
    )
    tabla_detalle = Tabla(
        "Detalle de novedades",
        [
            "Documento",
            "Empleado",
            "Sede",
            "Tipo",
            "Estado",
            "Fecha inicio",
            "Fecha fin",
            "Tiene soporte",
            "# soportes",
        ],
        detalle,
    )
    return Reporte(
        titulo="Reporte de novedades",
        subtitulo=_subtitulo(meta),
        meta=meta,
        tablas=[tabla_resumen, tabla_detalle],
    )


# ==========================================================================
# RECOLECCIÓN — exportación individual Habeas Data (Ley 1581)
# ==========================================================================
def _referencia_soporte(soporte) -> str:
    """Describe un soporte SIN volcar su contenido (dato sensible de salud)."""
    return (
        f"{soporte.nombre_original} "
        f"({soporte.tipo_mime}, {soporte.tamano_bytes} bytes; "
        f"{NOTA_SOPORTE_SENSIBLE})"
    )


def recolectar_habeas_data(empleado) -> Reporte:
    """Reúne TODOS los datos personales de un empleado en el sistema.

    Incluye: identidad, consentimientos, marcaciones, novedades (referenciando
    los soportes médicos SIN volcar el archivo) y resúmenes de jornada.
    """
    usuario = empleado.usuario

    # --- Identidad (clave/valor) ---
    identidad = Tabla(
        "Identidad",
        ["Campo", "Valor"],
        [
            ["Documento de identidad", empleado.documento_identidad],
            ["Nombres", empleado.nombres],
            ["Apellidos", empleado.apellidos],
            ["Sede", empleado.sede.nombre],
            ["Cargo", empleado.cargo],
            ["Fecha de ingreso", empleado.fecha_ingreso.isoformat()],
            ["Usuario", usuario.get_username()],
            ["Rol", usuario.get_rol_display()],
            ["Email", usuario.email or ""],
            ["Activo", _si_no(empleado.activo)],
        ],
    )

    # --- Consentimientos Habeas Data ---
    consentimientos = Tabla(
        "Consentimientos",
        ["Versión política", "Aceptado", "Fecha aceptación", "IP"],
        [
            [
                c.version_politica,
                _si_no(c.aceptado),
                _hora_local(c.fecha_aceptacion),
                c.ip_aceptacion,
            ]
            for c in empleado.consentimientos.all().order_by("-fecha_aceptacion")
        ],
    )

    # --- Marcaciones ---
    marcaciones = Tabla(
        "Marcaciones",
        ["Tipo", "Hora oficial", "Sede", "Fuera de sede", "Offline", "Corrección"],
        [
            [
                m.get_tipo_display(),
                _hora_local(_hora_oficial(m)),
                m.sede.nombre,
                _si_no(m.fuera_de_sede),
                _si_no(m.es_offline),
                _si_no(m.corrige_a_id is not None),
            ]
            for m in empleado.marcaciones.select_related("sede").order_by(
                "minuto_marcacion"
            )
        ],
    )

    # --- Novedades (soportes referenciados, NUNCA volcados) ---
    novedades_filas = []
    for n in empleado.novedades.prefetch_related("soportes").order_by("-created_at"):
        refs = "; ".join(_referencia_soporte(s) for s in n.soportes.all())
        novedades_filas.append(
            [
                n.get_tipo_display(),
                n.get_estado_display(),
                n.fecha_inicio.isoformat(),
                n.fecha_fin.isoformat() if n.fecha_fin else "",
                n.descripcion,
                refs,
            ]
        )
    novedades = Tabla(
        "Novedades",
        [
            "Tipo",
            "Estado",
            "Fecha inicio",
            "Fecha fin",
            "Descripción",
            "Soportes (referencia)",
        ],
        novedades_filas,
    )

    # --- Resúmenes de jornada ---
    resumenes = Tabla(
        "Resúmenes de jornada",
        [
            "Fecha",
            "H. Ordinarias",
            "Extra Diurnas",
            "Extra Nocturnas",
            "Recargo Nocturno",
            "Dominical/Festivo",
            "Tardanza (min)",
        ],
        [
            [
                r.fecha.isoformat(),
                _num(r.horas_ordinarias),
                _num(r.horas_extra_diurnas),
                _num(r.horas_extra_nocturnas),
                _num(r.horas_recargo_nocturno),
                _num(r.horas_dominical_festivo),
                r.tardanza_minutos,
            ]
            for r in empleado.resumenes_jornada.order_by("fecha")
        ],
    )

    meta = {
        "Titular": empleado.nombre_completo,
        "Documento": empleado.documento_identidad,
        "Generado": timezone.localtime().strftime("%Y-%m-%d %H:%M"),
        "Marco legal": "Ley 1581 de 2012 (Habeas Data)",
    }
    return Reporte(
        titulo="Exportación de datos personales (Habeas Data — Ley 1581)",
        subtitulo=f"Titular: {empleado.nombre_completo} ({empleado.documento_identidad})",
        meta=meta,
        tablas=[identidad, consentimientos, marcaciones, novedades, resumenes],
    )


# ==========================================================================
# ORQUESTACIÓN — render + AUDITORÍA obligatoria
# ==========================================================================
def exportar(
    *,
    reporte: Reporte,
    formato: str,
    actor,
    ip_origen: Optional[str],
    nombre_base: str,
    accion_meta: dict,
    entidad: str = "Reporte",
    entidad_id: Optional[Any] = None,
) -> ExportResult:
    """Renderiza el reporte y REGISTRA SIEMPRE la exportación en audit_log.

    Este es el único punto donde un reporte se materializa en bytes; por eso
    aquí se garantiza el rastro EXPORTACION_DATOS (acceso masivo a datos).
    """
    formato = (formato or "xlsx").lower()
    if formato == "xlsx":
        contenido = render_excel(reporte)
        content_type = XLSX_CONTENT_TYPE
        ext = "xlsx"
    elif formato == "pdf":
        contenido = render_pdf(reporte)
        content_type = PDF_CONTENT_TYPE
        ext = "pdf"
    else:
        raise ValueError(f"Formato no soportado: {formato!r} (use 'xlsx' o 'pdf').")

    filename = f"{nombre_base}.{ext}"

    metadata = dict(accion_meta)
    metadata.update(
        {
            "formato": formato,
            "filas": reporte.total_filas(),
            "bytes": len(contenido),
            "titulo": reporte.titulo,
        }
    )
    AuditService.registrar(
        accion=AuditLog.Accion.EXPORTACION_DATOS,
        entidad=entidad,
        entidad_id=str(entidad_id) if entidad_id is not None else None,
        actor_id=str(actor.pk) if actor is not None else None,
        actor_rol=getattr(actor, "rol", None),
        metadata=metadata,
        ip_origen=ip_origen,
    )
    return ExportResult(contenido=contenido, filename=filename, content_type=content_type)


def http_response(result: ExportResult) -> HttpResponse:
    """Envuelve un ExportResult en una descarga HTTP (attachment)."""
    resp = HttpResponse(result.contenido, content_type=result.content_type)
    resp["Content-Disposition"] = f'attachment; filename="{result.filename}"'
    return resp
