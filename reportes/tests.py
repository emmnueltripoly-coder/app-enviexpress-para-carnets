"""
Tests del Hito 8 (Reportes y exportación) — corren contra PostgreSQL real.

Cubren:
(a) Reporte de asistencia en Excel Y PDF: archivos válidos y no vacíos.
(b) Reporte de horas con las categorías legales tomadas de ResumenJornada.
(c) Toda exportación queda en audit_log con accion=EXPORTACION_DATOS.
(d) Exportación Habeas Data incluye los datos del empleado y NO vuelca el
    archivo médico (solo lo referencia).
(e) Un empleado NO puede exportar los datos de OTRO (403).
(f) RRHH SÍ puede exportar el Habeas Data de un empleado (200).
(g) El auditor puede generar un reporte de lectura y su acceso queda auditado.
"""

from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.utils import timezone
from rest_framework.test import APIClient

from auditoria.models import AuditLog
from marcacion.models import Marcacion
from novedades.models import Novedad, SoporteNovedad
from organizacion.models import Empresa, Sede
from personal.models import Empleado, Usuario
from reportes import services
from turnos.models import ResumenJornada

PDF_BYTES = b"%PDF-1.4 contenido confidencial de incapacidad medica"


@pytest.fixture(autouse=True)
def _private_media(settings, tmp_path):
    settings.PRIVATE_MEDIA_ROOT = str(tmp_path / "priv")


@pytest.fixture
def empresa():
    return Empresa.objects.create(nombre="Enviexpress", nit="900123456-7")


def _sede(empresa, nombre, ciudad):
    return Sede.objects.create(
        empresa=empresa,
        nombre=nombre,
        ciudad=ciudad,
        direccion="Calle 1",
        latitud="6.24",
        longitud="-75.58",
    )


def _empleado(sede, *, username, documento, rol=Usuario.Rol.EMPLEADO):
    usuario = Usuario.objects.create_user(
        username=username, password="ClaveSegura123", rol=rol
    )
    return Empleado.objects.create(
        usuario=usuario,
        documento_identidad=documento,
        nombres="Nom",
        apellidos="Ape",
        sede=sede,
        cargo="Cargo",
        fecha_ingreso=date(2026, 1, 1),
    )


def _marcacion(empleado, sede, *, tipo=Marcacion.Tipo.ENTRADA, cuando=None, fuera=False, offline=False):
    cuando = cuando or timezone.now()
    return Marcacion.objects.create(
        empleado=empleado,
        sede=sede,
        tipo=tipo,
        timestamp_qr=cuando,
        timestamp_servidor=cuando,
        latitud="6.24",
        longitud="-75.58",
        fuera_de_sede=fuera,
        distancia_metros=Decimal("12.50"),
        es_offline=offline,
        revisar_offline=offline,
    )


def _client(usuario):
    c = APIClient()
    c.force_authenticate(user=usuario)
    return c


def _hoy():
    return timezone.localdate()


# --------------------------------------------------------------------------
# (a) Reporte de asistencia en Excel y PDF — válidos y no vacíos
# --------------------------------------------------------------------------
@pytest.mark.django_db
def test_reporte_asistencia_excel_y_pdf_validos(empresa):
    sede = _sede(empresa, "Medellín", Sede.Ciudad.MEDELLIN)
    emp = _empleado(sede, username="cc.emp", documento="100")
    rrhh = Usuario.objects.create_user(username="rrhh", password="x", rol=Usuario.Rol.RRHH)
    _marcacion(emp, sede, tipo=Marcacion.Tipo.ENTRADA)
    _marcacion(emp, sede, tipo=Marcacion.Tipo.SALIDA, fuera=True)

    qs = Marcacion.objects.all()
    reporte = services.recolectar_asistencia(qs, meta={"Periodo": "hoy", "Sede": sede.nombre})

    xlsx = services.exportar(
        reporte=reporte, formato="xlsx", actor=rrhh, ip_origen="127.0.0.1",
        nombre_base="reporte_asistencia", accion_meta={"reporte": "asistencia"},
    )
    pdf = services.exportar(
        reporte=reporte, formato="pdf", actor=rrhh, ip_origen="127.0.0.1",
        nombre_base="reporte_asistencia", accion_meta={"reporte": "asistencia"},
    )

    # Excel: empieza por la firma ZIP (PK) y openpyxl lo puede reabrir.
    assert xlsx.contenido[:2] == b"PK"
    assert len(xlsx.contenido) > 0
    assert xlsx.content_type.endswith("spreadsheetml.sheet")
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(xlsx.contenido))
    assert "Asistencia" in wb.sheetnames

    # PDF: empieza por la firma %PDF y no está vacío.
    assert pdf.contenido[:4] == b"%PDF"
    assert len(pdf.contenido) > 0
    assert pdf.content_type == "application/pdf"


# --------------------------------------------------------------------------
# (b) Reporte de horas con las categorías legales desde ResumenJornada
# --------------------------------------------------------------------------
@pytest.mark.django_db
def test_reporte_horas_categorias_legales(empresa):
    sede = _sede(empresa, "Cali", Sede.Ciudad.CALI)
    emp = _empleado(sede, username="cc.emp", documento="100")
    ResumenJornada.objects.create(
        empleado=emp,
        fecha=_hoy(),
        horas_ordinarias=Decimal("8.00"),
        horas_extra_diurnas=Decimal("1.50"),
        horas_extra_nocturnas=Decimal("0.75"),
        horas_recargo_nocturno=Decimal("2.00"),
        horas_dominical_festivo=Decimal("3.25"),
        tardanza_minutos=5,
        calculado_at=timezone.now(),
    )

    reporte = services.recolectar_horas(
        ResumenJornada.objects.all(), meta={"Periodo": "hoy", "Sede": sede.nombre}
    )
    tabla = reporte.tablas[0]
    # Las 5 categorías legales están como columnas.
    for categoria in [
        "H. Ordinarias", "Extra Diurnas", "Extra Nocturnas",
        "Recargo Nocturno", "Dominical/Festivo",
    ]:
        assert categoria in tabla.columnas

    # Los valores del ResumenJornada aparecen en la fila de datos.
    fila = tabla.filas[0]
    assert 8.0 in fila and 1.5 in fila and 0.75 in fila and 2.0 in fila and 3.25 in fila
    # Hay una fila de TOTALES.
    assert any("TOTALES" in [str(c) for c in f] for f in tabla.filas)


# --------------------------------------------------------------------------
# (c) Toda exportación queda registrada en audit_log
# --------------------------------------------------------------------------
@pytest.mark.django_db
def test_exportacion_se_registra_en_audit_log(empresa):
    sede = _sede(empresa, "Bogotá", Sede.Ciudad.BOGOTA)
    emp = _empleado(sede, username="cc.emp", documento="100")
    rrhh = Usuario.objects.create_user(username="rrhh", password="x", rol=Usuario.Rol.RRHH)
    _marcacion(emp, sede)

    antes = AuditLog.objects.filter(accion=AuditLog.Accion.EXPORTACION_DATOS).count()
    reporte = services.recolectar_asistencia(Marcacion.objects.all(), meta={})
    services.exportar(
        reporte=reporte, formato="xlsx", actor=rrhh, ip_origen="10.0.0.9",
        nombre_base="reporte_asistencia",
        accion_meta={"reporte": "asistencia", "desde": "2026-06-01", "hasta": "2026-06-30"},
    )

    logs = AuditLog.objects.filter(accion=AuditLog.Accion.EXPORTACION_DATOS)
    assert logs.count() == antes + 1
    log = logs.order_by("-created_at").first()
    assert log.actor_id == str(rrhh.pk)
    assert log.actor_rol == Usuario.Rol.RRHH
    assert log.metadata["reporte"] == "asistencia"
    assert log.metadata["formato"] == "xlsx"
    assert log.ip_origen == "10.0.0.9"


# --------------------------------------------------------------------------
# (d) Habeas Data incluye los datos y NO vuelca el archivo médico
# --------------------------------------------------------------------------
@pytest.mark.django_db
def test_habeas_data_incluye_datos_sin_volcar_soporte(empresa):
    sede = _sede(empresa, "Medellín", Sede.Ciudad.MEDELLIN)
    emp = _empleado(sede, username="cc.emp", documento="DOC-555")
    novedad = Novedad.objects.create(
        empleado=emp, tipo=Novedad.Tipo.INCAPACIDAD,
        descripcion="Reposo por 3 días", fecha_inicio=_hoy(),
    )
    SoporteNovedad.objects.create(
        novedad=novedad,
        archivo=SimpleUploadedFile("incapacidad.pdf", PDF_BYTES, content_type="application/pdf"),
        nombre_original="incapacidad.pdf",
        tipo_mime="application/pdf",
        tamano_bytes=len(PDF_BYTES),
        subido_por=emp.usuario,
    )

    reporte = services.recolectar_habeas_data(emp)
    celdas = list(reporte.iter_celdas())
    texto = "\n".join(celdas)

    # SÍ incluye sus datos de identidad.
    assert "DOC-555" in texto
    # SÍ referencia el soporte (nombre del archivo) ...
    assert "incapacidad.pdf" in texto
    # ... con la nota de que el archivo NO se incluye.
    assert services.NOTA_SOPORTE_SENSIBLE in texto
    # Pero NUNCA vuelca el CONTENIDO del archivo médico.
    assert "contenido confidencial de incapacidad medica" not in texto
    for celda in celdas:
        assert PDF_BYTES.decode("latin-1") not in celda

    # El render tampoco contiene el contenido del archivo.
    pdf = services.render_pdf(reporte)
    xlsx = services.render_excel(reporte)
    assert b"contenido confidencial de incapacidad medica" not in pdf
    assert b"contenido confidencial de incapacidad medica" not in xlsx


# --------------------------------------------------------------------------
# (e) Un empleado NO puede exportar los datos de OTRO (403)
# --------------------------------------------------------------------------
@pytest.mark.django_db
def test_empleado_no_exporta_datos_de_otro(empresa):
    sede = _sede(empresa, "Medellín", Sede.Ciudad.MEDELLIN)
    emp_a = _empleado(sede, username="cc.a", documento="100")
    emp_b = _empleado(sede, username="cc.b", documento="200")

    # emp_a intenta exportar el expediente de emp_b -> 403.
    resp = _client(emp_a.usuario).get(f"/api/reportes/habeas-data/{emp_b.pk}/")
    assert resp.status_code == 403

    # emp_a SÍ puede exportar lo suyo por la misma ruta.
    propio = _client(emp_a.usuario).get(f"/api/reportes/habeas-data/{emp_a.pk}/")
    assert propio.status_code == 200

    # Y por el atajo /me/ también.
    yo = _client(emp_a.usuario).get("/api/reportes/habeas-data/me/")
    assert yo.status_code == 200


# --------------------------------------------------------------------------
# (f) RRHH SÍ puede exportar el Habeas Data de un empleado (200)
# --------------------------------------------------------------------------
@pytest.mark.django_db
def test_rrhh_exporta_habeas_data(empresa):
    sede = _sede(empresa, "Cali", Sede.Ciudad.CALI)
    emp = _empleado(sede, username="cc.emp", documento="100")
    rrhh = Usuario.objects.create_user(username="rrhh", password="x", rol=Usuario.Rol.RRHH)

    resp = _client(rrhh).get(f"/api/reportes/habeas-data/{emp.pk}/?formato=pdf")
    assert resp.status_code == 200
    assert resp["Content-Type"] == "application/pdf"
    assert b"".join(resp.streaming_content if hasattr(resp, "streaming_content") else [resp.content])[:4] == b"%PDF"
    # Quedó auditado contra el empleado titular.
    log = AuditLog.objects.filter(
        accion=AuditLog.Accion.EXPORTACION_DATOS, entidad="Empleado"
    ).order_by("-created_at").first()
    assert log is not None
    assert log.entidad_id == str(emp.pk)
    assert log.actor_rol == Usuario.Rol.RRHH


# --------------------------------------------------------------------------
# (g) El auditor genera un reporte de lectura y su acceso queda auditado
# --------------------------------------------------------------------------
@pytest.mark.django_db
def test_auditor_genera_reporte_y_queda_auditado(empresa):
    sede = _sede(empresa, "Barranquilla", Sede.Ciudad.BARRANQUILLA)
    emp = _empleado(sede, username="cc.emp", documento="100")
    auditor = Usuario.objects.create_user(username="auditor", password="x", rol=Usuario.Rol.AUDITOR)
    _marcacion(emp, sede)

    resp = _client(auditor).get("/api/reportes/asistencia/?formato=xlsx")
    assert resp.status_code == 200
    contenido = resp.content if hasattr(resp, "content") else b"".join(resp.streaming_content)
    assert contenido[:2] == b"PK"

    log = AuditLog.objects.filter(
        accion=AuditLog.Accion.EXPORTACION_DATOS, actor_id=str(auditor.pk)
    ).order_by("-created_at").first()
    assert log is not None
    assert log.actor_rol == Usuario.Rol.AUDITOR
    assert log.metadata["reporte"] == "asistencia"


# --------------------------------------------------------------------------
# (extra) Un empleado normal NO puede generar reportes agregados (403)
# --------------------------------------------------------------------------
@pytest.mark.django_db
def test_empleado_no_genera_reportes_agregados(empresa):
    sede = _sede(empresa, "Medellín", Sede.Ciudad.MEDELLIN)
    emp = _empleado(sede, username="cc.emp", documento="100")

    assert _client(emp.usuario).get("/api/reportes/asistencia/").status_code == 403
    assert _client(emp.usuario).get("/api/reportes/horas/").status_code == 403
    assert _client(emp.usuario).get("/api/reportes/novedades/").status_code == 403
    assert _client(emp.usuario).get("/api/reportes/excepciones/").status_code == 403
